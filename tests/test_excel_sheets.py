import os
import sys
import shutil
import tempfile
import unittest
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parents[1]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from openpyxl import Workbook, load_workbook

from src.helpers.excel import (
    DEFAULT_COMBINED_WORKBOOK_NAME,
    ExcelAppendWorkbook,
    ExcelStatusWorkbook,
    find_matching_sheet,
    format_cell_value,
    resolve_dataset,
)
from ckg_helper import MENU_OPTIONS, validate_excel_file


class TestExcelMultiSheet(unittest.TestCase):
    def setUp(self):
        self.temp_dir = Path(tempfile.mkdtemp())
        self.dataset_dir = self.temp_dir / "dataset"
        self.dataset_dir.mkdir(parents=True)

    def tearDown(self):
        shutil.rmtree(self.temp_dir)

    def test_find_matching_sheet_aliases(self):
        sheetnames = ["Pendaftaran Umum", "Konfirmasi Kehadiran", "Dewasa", "Anak Sekolah"]
        self.assertEqual(find_matching_sheet(sheetnames, "pendaftaran_umum"), "Pendaftaran Umum")
        self.assertEqual(find_matching_sheet(sheetnames, "konfirm_kehadiran"), "Konfirmasi Kehadiran")
        self.assertEqual(find_matching_sheet(sheetnames, "dewasa"), "Dewasa")
        self.assertIsNone(find_matching_sheet(sheetnames, "lansia"))

    def test_resolve_dataset_combined_vs_standalone(self):
        # 1. Initially no files exist
        target_path, sheet = resolve_dataset("dewasa", project_root=self.temp_dir)
        self.assertEqual(target_path, self.dataset_dir / "dewasa.xlsx")
        self.assertIsNone(sheet)

        # 2. Standalone file exists
        standalone = self.dataset_dir / "dewasa.xlsx"
        wb = Workbook()
        wb.active.append(["col1", "status"])
        wb.save(standalone)

        target_path, sheet = resolve_dataset("dewasa", project_root=self.temp_dir)
        self.assertEqual(target_path, standalone)
        self.assertIsNone(sheet)

        # 3. Combined workbook exists with matching sheet
        combined = self.dataset_dir / DEFAULT_COMBINED_WORKBOOK_NAME
        cwb = Workbook()
        cwb.active.title = "dewasa"
        cwb.active.append(["nik", "nama", "status"])
        cwb.save(combined)

        target_path, sheet = resolve_dataset("dewasa", project_root=self.temp_dir)
        self.assertEqual(target_path, combined)
        self.assertEqual(sheet, "dewasa")

    def test_status_update_on_specific_sheet(self):
        combined = self.dataset_dir / DEFAULT_COMBINED_WORKBOOK_NAME
        wb = Workbook()
        ws1 = wb.active
        ws1.title = "sheet_a"
        ws1.append(["id", "name", "status"])
        ws1.append(["1", "Alice", ""])

        ws2 = wb.create_sheet(title="sheet_b")
        ws2.append(["id", "name", "status"])
        ws2.append(["2", "Bob", "PENDING"])
        wb.save(combined)

        # Update sheet_a only
        excel_a = ExcelStatusWorkbook(combined, sheet_name="sheet_a")
        rows = excel_a.pending_rows()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["data"]["name"], "Alice")
        excel_a.update_status(2, "SUCCESS")

        # Verify sheet_a changed and sheet_b is untouched
        check_wb = load_workbook(combined)
        self.assertEqual(check_wb["sheet_a"].cell(row=2, column=3).value, "SUCCESS")
        self.assertEqual(check_wb["sheet_b"].cell(row=2, column=3).value, "PENDING")

    def test_append_row_to_dataset_combined_mode(self):
        combined = self.dataset_dir / DEFAULT_COMBINED_WORKBOOK_NAME
        wb = Workbook()
        ws_reg = wb.active
        ws_reg.title = "pendaftaran_umum"
        ws_reg.append(["nik", "nama_lengkap", "status"])
        ws_reg.append(["12345", "Budi", ""])

        ws_att = wb.create_sheet(title="konfirm_kehadiran")
        ws_att.append(["nama_lengkap", "tgl_pemeriksaan", "status"])
        wb.save(combined)

        # Open pendaftaran_umum and process row
        excel = ExcelStatusWorkbook(combined, sheet_name="pendaftaran_umum")
        excel.update_status(2, "SUCCESS")
        excel.append_row_to_dataset("konfirm_kehadiran", {
            "nama_lengkap": "Budi",
            "tgl_pemeriksaan": "2026-09-20",
        })

        # Verify both sheets in the file were updated without clobbering
        check_wb = load_workbook(combined)
        self.assertEqual(check_wb["pendaftaran_umum"].cell(row=2, column=3).value, "SUCCESS")
        self.assertEqual(check_wb["konfirm_kehadiran"].cell(row=2, column=1).value, "Budi")
        self.assertEqual(check_wb["konfirm_kehadiran"].cell(row=2, column=2).value, "2026-09-20")

    def test_append_row_to_dataset_skips_partially_populated_rows(self):
        combined = self.dataset_dir / DEFAULT_COMBINED_WORKBOOK_NAME
        wb = Workbook()
        ws_reg = wb.active
        ws_reg.title = "pendaftaran_umum"
        ws_reg.append(["nama_lengkap", "status"])

        ws_anak = wb.create_sheet(title="anak")
        ws_anak.append(["tgl_pemeriksaan", "nama"])
        ws_anak.cell(row=2, column=2, value="Existing")
        wb.save(combined)

        excel = ExcelStatusWorkbook(combined, sheet_name="pendaftaran_umum")
        excel.append_row_to_dataset("anak", {"nama": "Budi"})

        check_wb = load_workbook(combined)
        self.assertEqual(check_wb["anak"].cell(row=2, column=2).value, "Existing")
        self.assertEqual(check_wb["anak"].cell(row=3, column=2).value, "Budi")

    def test_validate_excel_file_with_combined_and_standalone(self):
        # Combined workbook with all 9 sheets
        combined = self.dataset_dir / DEFAULT_COMBINED_WORKBOOK_NAME
        wb = Workbook()
        for idx, (key, opt) in enumerate(MENU_OPTIONS.items()):
            sheet_title = opt.get("sheet", opt.get("dataset_key"))
            if idx == 0:
                wb.active.title = sheet_title
            else:
                wb.create_sheet(title=sheet_title)
        wb.save(combined)

        for key, opt in MENU_OPTIONS.items():
            valid = validate_excel_file(self.temp_dir, opt)
            self.assertTrue(valid, f"Menu {key} validation should succeed with combined workbook")


if __name__ == "__main__":
    unittest.main()
