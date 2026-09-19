#!/usr/bin/env python3
"""
Combines all standalone Excel files in dataset/ into a single multi-sheet dataset/ckg_data.xlsx workbook.
"""
from pathlib import Path
from openpyxl import Workbook, load_workbook

FILES_TO_SHEETS = [
    ("pendaftaran_umum.xlsx", "pendaftaran_umum"),
    ("konfirm_kehadiran.xlsx", "konfirm_kehadiran"),
    ("anak.xlsx", "anak"),
    ("remaja.xlsx", "remaja"),
    ("dewasa.xlsx", "dewasa"),
    ("lansia.xlsx", "lansia"),
    ("pendaftaran_sekolah.xlsx", "pendaftaran_sekolah"),
    ("konfirm_kehadiran_sekolah.xlsx", "konfirm_kehadiran_sekolah"),
    ("pelayanan_sekolah.xlsx", "pelayanan_sekolah"),
]


def combine_excel(dataset_dir: Path | None = None, output_filename: str = "ckg_data.xlsx") -> Path | None:
    if dataset_dir is None:
        dataset_dir = Path(__file__).resolve().parents[1] / "dataset"

    out_path = dataset_dir / output_filename
    combined_wb = Workbook()
    default_sheet = combined_wb.active

    sheets_added = 0
    for filename, sheet_name in FILES_TO_SHEETS:
        file_path = dataset_dir / filename
        if not file_path.exists():
            print(f"Skipping {filename} (file not found)")
            continue

        src_wb = load_workbook(file_path, data_only=False)
        src_sheet = src_wb.active

        dest_sheet = combined_wb.create_sheet(title=sheet_name)
        for row in src_sheet.iter_rows(values_only=False):
            dest_sheet.append([cell.value for cell in row])

        print(f"✓ Added sheet '{sheet_name}' from {filename} ({src_sheet.max_row} rows, {src_sheet.max_column} cols)")
        sheets_added += 1

    if default_sheet in combined_wb.worksheets:
        combined_wb.remove(default_sheet)

    if sheets_added > 0:
        combined_wb.save(out_path)
        print(f"\nBerhasil menggabungkan {sheets_added} sheet ke: {out_path}")
        return out_path
    else:
        print("\nTidak ada file Excel yang ditemukan untuk digabungkan.")
        return None


if __name__ == "__main__":
    combine_excel()
